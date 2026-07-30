#!/usr/bin/env python3
"""GAINS MAP backend.

Serves APP.html and a small JSON API. All data lives in gains.xlsx:
configuration (profile, groups, exercises, levels), muscle XP and the
daily workout log. Edit the Excel file to change exercises/levels/profile;
the frontend loads everything from here.
"""
import os
from datetime import date

from flask import Flask, jsonify, request, send_from_directory
from openpyxl import Workbook, load_workbook

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
XLSX_PATH = os.path.join(BASE_DIR, "gains.xlsx")

app = Flask(__name__)

# ---------------------------------------------------------------- defaults
DEFAULT_PROFILE = [
    ("sex", "♂"), ("age", 36), ("height_cm", 180),
    ("weight_kg", 70), ("bmi", 22), ("bodyfat_pct", 19),
]

DEFAULT_GROUPS = [
    # key, label, hint
    ("abs", "Abs", ""),
    ("chest", "Chest", "push forward"),
    ("shoulders", "Shoulders", "push up"),
    ("triceps", "Triceps", "push down"),
    ("back", "Back", "pull down / pull back"),
    ("biceps", "Biceps", "pull up"),
    ("legs", "Legs", ""),
    ("mobility", "Mobility", ""),
]

DEFAULT_XP = [
    # group, starting xp, weighted ('yes' = uses weights, kg follows the level)
    ("abs", 40, ""), ("chest", 40, "yes"), ("shoulders", 40, "yes"),
    ("triceps", 24, "yes"), ("back", 40, "yes"), ("biceps", 40, "yes"),
    ("legs", 24, "yes"), ("mobility", 30, ""),
]

DEFAULT_EXERCISES = [
    # id, group, label, impact ("region:weight;region:weight")
    ("plank", "abs", "Isometric (plank gambe/braccia piegate)", "abs_upper:1"),
    ("legraise", "abs", "Crunch (panca alza gambe)", "abs_lower:1"),
    ("oblique", "abs", "Oblique", "obliques:1"),
    ("chestpress45", "chest", "ChestPress inclinata 45", "chest:1;triceps:0.45"),
    ("chestpressflat", "chest", "ChestPress flat", "chest:1;triceps:0.45"),
    ("fly", "chest", "PectoralFly", "chest:1;triceps:0.45"),
    ("shoulderpress", "shoulders", "ShoulderPress (inclinata 85)", "shoulders:1;triceps:0.45"),
    ("ropepushdown", "triceps", "RopePushDown", "triceps:1"),
    ("benchdips", "triceps", "BenchDips", "triceps:1"),
    ("latpulldown", "back", "LatPullDown", "back:1;biceps:0.45"),
    ("seatedrow", "back", "SeatedRow", "back:1;biceps:0.45"),
    ("curls", "biceps", "DumbbellStandingCurls", "biceps:1"),
    ("squat", "legs", "Squat (con pesi)", "upperlegs_front:1"),
    ("affondi", "legs", "Affondi (verso dietro)", "upperlegs_back:1"),
    ("calfraises", "legs", "CalfRaises (con pesi)", "calves_back:1"),
    ("stretching", "mobility", "Stretching", "head:1"),
    ("running", "mobility", "Running", "head:1"),
    ("sport", "mobility", "Any sport", "head:1"),
]

DEFAULT_LEVELS = [
    # min, max, square, name, color, kg (weight unlocked at this level: 5kg + 5 per level)
    (0, 9, "⬛", "Disgusting", "#3a3a3f", 5),
    (10, 19, "⬜", "Pathetic", "#b9b9c0", 10),
    (20, 29, "🟥", "Fragile", "#e0393f", 15),
    (30, 39, "🟥", "Weak", "#e0393f", 20),
    (40, 49, "🟧", "Yamcha", "#f08c2e", 25),
    (50, 59, "🟧", "Krillin", "#f08c2e", 30),
    (60, 69, "🟨", "Warrior", "#f2c14e", 35),
    (70, 79, "🟨", "Ninja", "#f2c14e", 40),
    (80, 89, "🟩", "Piccolo", "#4caf6e", 45),
    (90, 99, "🟪", "Toji", "#a86ee0", 50),
    (100, 999, "🟦", "Vegeta", "#4a8fe0", 55),
]

DEFAULT_PLAN = [
    # day, groups scheduled ("" = rest). Skipping a day costs 1 XP per group.
    ("Mon", "chest;shoulders;triceps"),
    ("Tue", "legs;back;biceps"),
    ("Wed", "abs;mobility"),
    ("Thu", "chest;shoulders;triceps"),
    ("Fri", "legs;back;biceps"),
    ("Sat", "abs;mobility"),
    ("Sun", ""),
]

SHEETS = {
    "Profile": ["key", "value"],
    "Groups": ["key", "label", "hint"],
    "XP": ["group", "xp", "weighted"],
    "Exercises": ["id", "group", "label", "impact"],
    "Levels": ["min", "max", "square", "name", "color", "kg"],
    "Workouts": ["date", "exercise_id"],
    "Plan": ["day", "groups"],
    "Skips": ["date"],
}


def create_default_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    for name, headers in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(headers)
    for row in DEFAULT_PROFILE:
        wb["Profile"].append(row)
    for row in DEFAULT_GROUPS:
        wb["Groups"].append(row)
    for row in DEFAULT_XP:
        wb["XP"].append(row)
    for row in DEFAULT_EXERCISES:
        wb["Exercises"].append(row)
    for row in DEFAULT_LEVELS:
        wb["Levels"].append(row)
    for row in DEFAULT_PLAN:
        wb["Plan"].append(row)
    wb.save(XLSX_PATH)
    return wb


def get_workbook():
    if not os.path.exists(XLSX_PATH):
        return create_default_workbook()
    wb = load_workbook(XLSX_PATH)
    # add sheets introduced after the file was first created
    changed = False
    for name, headers in SHEETS.items():
        if name not in wb.sheetnames:
            ws = wb.create_sheet(name)
            ws.append(headers)
            if name == "Plan":
                for row in DEFAULT_PLAN:
                    ws.append(row)
            changed = True
    if changed:
        wb.save(XLSX_PATH)
    return wb


def sheet_rows(ws):
    """Rows of a sheet as list of dicts keyed by the header row."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h) for h in rows[0]]
    out = []
    for row in rows[1:]:
        if all(c is None for c in row):
            continue
        out.append(dict(zip(headers, row)))
    return out


def parse_impact(text):
    """'chest:1;triceps:0.45' -> {'chest': 1.0, 'triceps': 0.45}"""
    impact = {}
    for part in str(text or "").split(";"):
        part = part.strip()
        if not part:
            continue
        region, _, weight = part.partition(":")
        try:
            impact[region.strip()] = float(weight)
        except ValueError:
            pass
    return impact


def num(v):
    """Number from an Excel cell; whole floats become ints (25.0 -> 25)."""
    f = float(v)
    return int(f) if f == int(f) else f


def iso(d):
    """Normalize a Workouts date cell (datetime or string) to YYYY-MM-DD."""
    if hasattr(d, "strftime"):
        return d.strftime("%Y-%m-%d")
    return str(d)[:10]


def build_state(wb):
    groups = [
        {"key": g["key"], "label": g["label"], "hint": g.get("hint") or ""}
        for g in sheet_rows(wb["Groups"])
    ]
    exercises = [
        {"id": e["id"], "group": e["group"], "label": e["label"],
         "impact": parse_impact(e["impact"])}
        for e in sheet_rows(wb["Exercises"])
    ]

    levels = [
        {"min": int(l["min"]), "max": int(l["max"]), "sq": l["square"],
         "name": l["name"], "color": l["color"],
         "kg": num(l["kg"]) if l.get("kg") not in ("", None) else None}
        for l in sheet_rows(wb["Levels"])
    ]

    def level_for(v):
        for l in levels:
            if l["min"] <= v <= l["max"]:
                return l
        return levels[0]

    plan = []
    for p in sheet_rows(wb["Plan"]):
        plan.append({"day": str(p["day"]),
                     "groups": [g.strip() for g in str(p.get("groups") or "").split(";") if g.strip()]})
    skips = sorted({iso(s["date"]) for s in sheet_rows(wb["Skips"])}, reverse=True)

    # current XP = starting XP (XP sheet) + 1 per exercise ever logged
    # - 1 per scheduled group on every skipped day. Workouts/Skips sheets
    # are the source of truth. Weighted groups train at their level's kg.
    group_of = {str(e["id"]): str(e["group"]) for e in sheet_rows(wb["Exercises"])}
    earned = {}
    for w in sheet_rows(wb["Workouts"]):
        g = group_of.get(str(w["exercise_id"]))
        if g:
            earned[g] = earned.get(g, 0) + 1
    lost = {}
    day_names = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
    plan_by_day = {p["day"]: p["groups"] for p in plan}
    for s in skips:
        try:
            weekday = day_names[date(*[int(x) for x in s.split("-")]).weekday()]
        except (ValueError, TypeError):
            continue
        for g in plan_by_day.get(weekday, []):
            lost[g] = lost.get(g, 0) + 1
    xp = {}
    for r in sheet_rows(wb["XP"]):
        group = str(r["group"])
        value = max(0, min(100, int(r["xp"] or 0) + earned.get(group, 0) - lost.get(group, 0)))
        weighted = str(r.get("weighted") or "").strip().lower() in ("yes", "y", "1", "true")
        xp[group] = {"xp": value,
                     "kg": level_for(value)["kg"] if weighted else None}
    profile = {str(p["key"]): p["value"] for p in sheet_rows(wb["Profile"])}

    entries_by_date = {}
    for w in sheet_rows(wb["Workouts"]):
        d = iso(w["date"])
        entries_by_date.setdefault(d, []).append(str(w["exercise_id"]))
    entries = [{"date": d, "exIds": ids} for d, ids in
               sorted(entries_by_date.items(), reverse=True)]

    return {"groups": groups, "xp": xp, "exercises": exercises,
            "levels": levels, "profile": profile, "entries": entries,
            "plan": plan, "skips": skips}


# ---------------------------------------------------------------- routes
@app.route("/")
def index():
    return send_from_directory(BASE_DIR, "APP.html")


@app.route("/api/state")
def api_state():
    return jsonify(build_state(get_workbook()))


@app.route("/api/workout", methods=["POST"])
def api_workout():
    payload = request.get_json(force=True) or {}
    day = str(payload.get("date") or date.today().isoformat())[:10]
    ex_ids = [str(x) for x in payload.get("exIds") or []]
    if not ex_ids:
        return jsonify({"error": "no exercises given"}), 400

    wb = get_workbook()
    known = {str(e["id"]): str(e["group"]) for e in sheet_rows(wb["Exercises"])}
    bad = [x for x in ex_ids if x not in known]
    if bad:
        return jsonify({"error": "unknown exercise ids: %s" % ", ".join(bad)}), 400

    ws = wb["Workouts"]
    for ex_id in ex_ids:
        ws.append([day, ex_id])

    wb.save(XLSX_PATH)
    return jsonify(build_state(wb))


@app.route("/api/skip", methods=["POST"])
def api_skip():
    """Toggle a skipped-training day: log it (costs 1 XP per scheduled
    group) or un-log it if already marked."""
    payload = request.get_json(force=True) or {}
    day = str(payload.get("date") or "")[:10]
    try:
        date(*[int(x) for x in day.split("-")])
    except (ValueError, TypeError):
        return jsonify({"error": "invalid date"}), 400

    wb = get_workbook()
    ws = wb["Skips"]
    existing = [i for i in range(2, ws.max_row + 1)
                if iso(ws.cell(row=i, column=1).value) == day]
    if existing:
        for i in reversed(existing):
            ws.delete_rows(i)
    else:
        ws.append([day])
    wb.save(XLSX_PATH)
    return jsonify(build_state(wb))


if __name__ == "__main__":
    get_workbook()  # create gains.xlsx on first run
    print("GAINS MAP running on http://localhost:5001  (data: %s)" % XLSX_PATH)
    app.run(host="127.0.0.1", port=5001, debug=False)
